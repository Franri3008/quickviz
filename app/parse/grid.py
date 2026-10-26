import csv
import io
import json
import re
import zipfile
from xml.etree import ElementTree as ET

from app.parse.values import MAX_CELL, MAX_COLS, MAX_ROWS, ParseError, cell, nonempty, to_number


class Grid(list):
    sheet = None
    sep = ""
    fmt = ""


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("Could not read the text encoding of this file")


def _sniff(lines):
    tabs = [ln.count("\t") for ln in lines]
    if tabs and tabs[0] > 0 and sum(1 for c in tabs if c == tabs[0]) >= 0.9 * len(tabs):
        return "\t"
    best, best_score = ",", -1
    for sep in [",", ";", "\t", "|"]:
        counts = [ln.count(sep) for ln in lines]
        common = max(set(counts), key=counts.count)
        if common == 0:
            continue
        score = sum(1 for c in counts if c == common) * 10 + common
        if score > best_score:
            best, best_score = sep, score
    return best


def _text_grid(raw: bytes):
    text = _decode(raw).replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln for ln in text.split("\n") if ln.strip()][:50]
    if not lines:
        raise ParseError("The file is empty")
    best = _sniff(lines)
    reader = csv.reader(io.StringIO(text), delimiter=best)
    grid = []
    for i, row in enumerate(reader):
        if i > MAX_ROWS + 60:
            break
        grid.append([c.strip()[:MAX_CELL] for c in row[:MAX_COLS]])
    if best == "\t" and grid and grid[0] and "\\" in grid[0][0] and "," in grid[0][0]:
        k = grid[0][0].count(",")
        out = []
        for r in grid:
            if r and r[0].count(",") == k:
                parts = [p.strip() for p in r[0].split(",")]
                r = parts + r[1:]
            out.append(r)
        out[0][k] = out[0][k].split("\\")[0].strip()
        grid = out
    return grid, best


def _merges(raw: bytes):
    out = {}
    try:
        z = zipfile.ZipFile(io.BytesIO(raw))
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
              "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        wbx = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target") for r in rels}
        for sh in wbx.find("m:sheets", ns):
            rid = sh.get(f"{{{ns['r']}}}id")
            t = targets.get(rid, "")
            path = t.lstrip("/") if t.startswith("/") else "xl/" + t
            try:
                xml = z.read(path).decode("utf-8", "ignore")
            except KeyError:
                continue
            refs = re.findall(r'<(?:\w+:)?mergeCell [^>]*ref="([A-Z]+\d+:[A-Z]+\d+)"', xml)
            if refs:
                out[sh.get("name")] = refs
    except Exception:
        return {}
    return out


def _addr(a):
    m = re.match(r"([A-Z]+)(\d+)", a)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)) - 1, col - 1


def _fill_merges(grid, refs):
    for ref in refs:
        a, b = ref.split(":")
        r0, c0 = _addr(a)
        r1, c1 = _addr(b)
        if r0 >= len(grid) or c0 >= MAX_COLS:
            continue
        v = grid[r0][c0] if c0 < len(grid[r0]) else ""
        if v == "":
            continue
        for r in range(r0, min(r1, len(grid) - 1) + 1):
            row = grid[r]
            for c in range(c0, min(c1, MAX_COLS - 1) + 1):
                while len(row) <= c:
                    row.append("")
                if row[c] == "":
                    row[c] = v


def _xlsx_sheets(raw: bytes):
    try:
        import openpyxl
    except ImportError as e:
        raise ParseError("Excel support is not installed on the server") from e
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ParseError("Could not open this Excel file. Save it as .xlsx and try again") from e
    out = {}
    for ws in wb.worksheets:
        grid = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS + 60:
                break
            grid.append([cell(c) for c in row[:MAX_COLS]])
        out[ws.title] = grid
    wb.close()
    for name, refs in _merges(raw).items():
        if name in out:
            _fill_merges(out[name], refs)
    return out


def _xls_sheets(raw: bytes):
    try:
        import xlrd
    except ImportError as e:
        raise ParseError("Old .xls files are not supported. Save it as .xlsx and try again") from e
    book = xlrd.open_workbook(file_contents=raw)
    out = {}
    for sh in book.sheets():
        grid = []
        for r in range(min(sh.nrows, MAX_ROWS + 60)):
            row = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cl = sh.cell(r, c)
                if cl.ctype == xlrd.XL_CELL_DATE:
                    row.append(cell(xlrd.xldate.xldate_as_datetime(cl.value, book.datemode)))
                else:
                    row.append(cell(cl.value))
            grid.append(row)
        for r0, r1, c0, c1 in getattr(sh, "merged_cells", []) or []:
            if r0 < len(grid):
                v = grid[r0][c0] if c0 < len(grid[r0]) else ""
                for r in range(r0, min(r1, len(grid))):
                    for c in range(c0, min(c1, len(grid[r]))):
                        if grid[r][c] == "":
                            grid[r][c] = v
        out[sh.name] = grid
    return out


def _jsonstat(j):
    if j.get("class") == "collection" or "dimension" not in j:
        return None
    ids = j.get("id") or list(j["dimension"])
    sizes = j.get("size") or []
    dims = []
    for k, d in zip(ids, sizes):
        dim = j["dimension"][k]
        cat = dim.get("category", {})
        idx = cat.get("index")
        if isinstance(idx, list):
            keys = idx
        elif isinstance(idx, dict):
            keys = sorted(idx, key=idx.get)
        else:
            keys = list((cat.get("label") or {}).keys())
        labels = cat.get("label") or {}
        dims.append((dim.get("label") or k, [labels.get(x, x) for x in keys][:d] or [""] * d))
    vals = j.get("value", [])
    total = 1
    for s in sizes:
        total *= s
    if isinstance(vals, dict):
        vals = [vals.get(str(i)) for i in range(total)]
    grid = [[name for name, _ in dims] + ["Value"]]
    for i in range(min(total, MAX_ROWS)):
        rest, row = i, []
        for _, labels in reversed(dims):
            row.append(labels[rest % len(labels)])
            rest //= len(labels)
        row.reverse()
        v = vals[i] if i < len(vals) else None
        row.append(cell(v))
        grid.append([cell(c) for c in row])
    return grid


def _json_grid(raw: bytes):
    try:
        j = json.loads(_decode(raw))
    except ValueError as e:
        raise ParseError("Could not read this JSON file") from e
    if isinstance(j, dict) and "dimension" in j and "value" in j:
        g = _jsonstat(j)
        if g:
            return g
    if isinstance(j, dict):
        arrays = [v for v in j.values() if isinstance(v, list)]
        if len(arrays) == 1:
            j = arrays[0]
        elif arrays and len(arrays) == len(j):
            n = max(len(a) for a in arrays)
            j = [{k: (v[i] if i < len(v) else None) for k, v in j.items()} for i in range(n)]
    if not isinstance(j, list) or not j or not isinstance(j[0], dict):
        raise ParseError("JSON must be a list of records")
    names = []
    for r in j[:MAX_ROWS]:
        if isinstance(r, dict):
            for k in r:
                if k not in names:
                    names.append(k)
    names = names[:MAX_COLS]
    grid = [names]
    for r in j[:MAX_ROWS]:
        if isinstance(r, dict):
            grid.append([cell(json.dumps(r.get(n)) if isinstance(r.get(n), (dict, list)) else r.get(n)) for n in names])
    return grid


def _sheet_score(grid):
    rows = [r for r in grid[:300] if nonempty(r)]
    if len(rows) < 2:
        return 0
    wide = sum(1 for r in rows if nonempty(r) >= 2)
    nums = sum(1 for r in rows for c in r if c != "" and to_number(c, False) is not None)
    if wide == 0:
        return 1
    return 3 if nums else 2


def _pick(books, sheet):
    names = list(books)
    if sheet and sheet in books:
        return sheet
    scores = {s: _sheet_score(books[s]) for s in names}
    top = max(scores.values(), default=0)
    if top == 0:
        raise ParseError("No sheet in this workbook has data")
    return next(s for s in names if scores[s] == top)


def read_grid(raw: bytes, name: str, sheet: str | None = None):
    if not raw:
        raise ParseError("The file is empty")
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    books = None
    if ext in ("xlsx", "xlsm") or raw[:4] == b"PK\x03\x04":
        books = _xlsx_sheets(raw)
        fmt = "xlsx"
    elif ext == "xls" or raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        books = _xls_sheets(raw)
        fmt = "xls"
    if books is not None:
        chosen = _pick(books, sheet)
        g = Grid(books[chosen])
        g.sheet, g.fmt = chosen, fmt
        return g, list(books)
    if ext == "json" or raw.lstrip()[:1] in (b"[", b"{"):
        g = Grid(_json_grid(raw))
        g.fmt = "json"
        return g, []
    rows, sep = _text_grid(raw)
    g = Grid(rows)
    g.sep, g.fmt = sep, "text"
    return g, []

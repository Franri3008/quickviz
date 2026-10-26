import csv
import datetime as dt
import io
import json
import random
import shutil
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]

EU = ["Belgium", "Bulgaria", "Czechia", "Denmark", "Germany", "Estonia", "Ireland", "Greece", "Spain", "France",
      "Croatia", "Italy", "Cyprus", "Latvia", "Lithuania", "Luxembourg", "Hungary", "Malta", "Netherlands", "Austria",
      "Poland", "Portugal", "Romania", "Slovenia", "Slovakia", "Finland", "Sweden"]
ISO2 = ["BE", "BG", "CZ", "DK", "DE", "EE", "IE", "EL", "ES", "FR", "HR", "IT", "CY", "LV", "LT", "LU", "HU", "MT", "NL",
        "AT", "PL", "PT", "RO", "SI", "SK", "FI", "SE"]
ISO3 = ["BEL", "BGR", "CZE", "DNK", "DEU", "EST", "IRL", "GRC", "ESP", "FRA", "HRV", "ITA", "CYP", "LVA", "LTU", "LUX",
        "HUN", "MLT", "NLD", "AUT", "POL", "PRT", "ROU", "SVN", "SVK", "FIN", "SWE"]


def expect(name, columns, n_rows, rows, about, aggregates=None, not_aggregates=None, sheet=None, needs_ai=False):
    data = {"about": about, "sheet": sheet, "needs_ai": needs_ai, "columns": columns, "n_rows": n_rows, "rows": rows}
    if aggregates is not None:
        data["aggregates"] = aggregates
    if not_aggregates is not None:
        data["not_aggregates"] = not_aggregates
    (HERE / f"{name}.expected.json").write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def write_text(name, text, encoding="utf-8", newline="\n"):
    (HERE / name).write_bytes(text.replace("\n", newline).encode(encoding))


def csv_text(rows, sep=","):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=sep, lineterminator="\n")
    for r in rows:
        w.writerow(r)
    return buf.getvalue()


def save_wb(wb, name):
    wb.save(HERE / name)


def stress():
    shutil.copy(ROOT / "samples" / "STRESS_eurostat_style_ai_use_by_size.xlsx", HERE / "eurostat_stress_three_tables.xlsx")
    lab = "Size classes in number of persons employed"
    expect("eurostat_stress_three_tables.xlsx",
           [[lab, "category"], ["GEO", "category"], ["Year", "time"], ["Value", "number"], ["flag", "category"]], 371,
           [{lab: "250 persons employed or more", "GEO": "Belgium", "Year": "2023", "Value": 43.07, "flag": "b"},
            {lab: "From 10 to 249 persons employed", "GEO": "Ireland", "Year": "2025", "Value": 14.12, "flag": "p"},
            {lab: "10 persons employed or more", "GEO": "Türkiye", "Year": "2021", "Value": 1.51, "flag": ""},
            {lab: "250 persons employed or more", "GEO": "Malta", "Year": "2023", "Value": 39.61, "flag": "u"}],
           "The stress file. Eurostat XLSX with three size-class tables stacked, flag columns, legend footer",
           aggregates=["European Union - 27 countries (from 2020)", "Euro area – 20 countries (from 2023)"],
           not_aggregates=["Belgium", "Malta"])


def eurostat_single():
    rnd = random.Random(2)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet 1"
    for r in [["Enterprises using cloud computing services [isoc_cicce_use__custom_1]"], ["Open product page"],
              ["Open in Data Browser"], [], ["Description:", "Mock layout of a single-table Eurostat download"],
              ["Last update", "02/09/2026 23:00"], ["Extracted on", "12/10/2026 09:12"], ["Source of data", "Eurostat"], [],
              ["Time frequency", "Annual"], ["Information society indicator", "Buy cloud computing services"],
              ["Size classes in number of persons employed", "10 persons employed or more"],
              ["Unit of measure", "Percentage of enterprises"], [], ["TIME", "2020", "", "2021", "", "2023", "", "2025", ""],
              ["GEO (Labels)", "", "", "", "", "", "", "", ""]]:
        ws.append(r)
    geos = ["European Union - 27 countries (from 2020)"] + EU[:12]
    rows, n = [], 0
    for g in geos:
        vals = []
        for y in range(4):
            if g == "Greece" and y == 0:
                vals += [":", ""]
            elif g == "Spain" and y == 3:
                vals += [round(rnd.uniform(20, 60), 1), "p"]
                n += 1
            else:
                vals += [round(rnd.uniform(20, 60), 1), "b" if g == "France" and y == 1 else ""]
                n += 1
        ws.append([g] + vals)
        rows.append([g] + vals)
    for r in [[], [], ["Special value"], [":", "not available"], ["Available flags:"], ["b", "break in time series"],
              ["p", "provisional"]]:
        ws.append(r)
    save_wb(wb, "eurostat_single_table.xlsx")
    fr = rows[geos.index("France")]
    es = rows[geos.index("Spain")]
    expect("eurostat_single_table.xlsx", [["GEO", "category"], ["Year", "time"], ["Value", "number"], ["flag", "category"]], n,
           [{"GEO": "France", "Year": "2021", "Value": fr[3], "flag": "b"},
            {"GEO": "Spain", "Year": "2025", "Value": es[7], "flag": "p"},
            {"GEO": "Belgium", "Year": "2020", "Value": rows[1][1], "flag": ""}],
           "Eurostat XLSX, one table, flag columns, metadata block on top, legend below",
           aggregates=["European Union - 27 countries (from 2020)"], not_aggregates=["Belgium"])


def eurostat_noflags():
    rnd = random.Random(3)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet 1"
    for r in [["Gross domestic expenditure on R&D by sector [rd_e_gerdtot]"], [], ["Last update", "11/09/2026"],
              ["Unit of measure", "Percentage of gross domestic product (GDP)"], ["Sector of performance", "All sectors"], [],
              ["TIME", 2019, 2020, 2021, 2022, 2023], ["GEO (Labels)"]]:
        ws.append(r)
    out, n = [], 0
    for g in ["EU27_2020"] + EU[:10]:
        vals = []
        for y in range(5):
            if g == "Bulgaria" and y == 4:
                vals.append(":")
            else:
                vals.append(round(rnd.uniform(0.5, 3.5), 2))
                n += 1
        ws.append([g] + vals)
        out.append([g] + vals)
    ws.append([])
    ws.append(["Special value"])
    ws.append([":", "not available"])
    save_wb(wb, "eurostat_noflags.xlsx")
    expect("eurostat_noflags.xlsx", [["GEO", "category"], ["Year", "time"], ["Value", "number"]], n,
           [{"GEO": "Germany", "Year": "2021", "Value": out[5][3]}, {"GEO": "EU27_2020", "Year": "2019", "Value": out[0][1]}],
           "Eurostat XLSX without flag columns, numeric year headers, single GEO row",
           aggregates=["EU27_2020"], not_aggregates=["Germany"])


def eurostat_tsv():
    rnd = random.Random(4)
    years = ["2019", "2020", "2021", "2022", "2023"]
    lines = ["freq,unit,geo\\TIME_PERIOD\t" + "\t".join(y + " " for y in years)]
    n, spots = 0, []
    for g in ["AT", "BE", "BG", "DE", "EA20", "EU27_2020", "FR", "IT", "NL"]:
        cells = []
        for k, y in enumerate(years):
            v = round(rnd.uniform(5, 40), 1)
            if g == "BG" and k == 0:
                cells.append(": ")
            elif g == "IT" and k == 2:
                cells.append(": c")
            elif g == "FR" and k == 4:
                cells.append(f"{v} p")
                n += 1
                spots.append({"geo": "FR", "Year": y, "Value": v, "flag": "p"})
            elif g == "DE" and k == 1:
                cells.append(f"{v} b")
                n += 1
                spots.append({"geo": "DE", "Year": y, "Value": v, "flag": "b"})
            elif g == "NL" and k == 3:
                cells.append(f"{v} bp")
                n += 1
                spots.append({"geo": "NL", "Year": y, "Value": v, "flag": "bp"})
            else:
                cells.append(f"{v} ")
                n += 1
                if g == "AT" and k == 0:
                    spots.append({"geo": "AT", "Year": y, "Value": v, "flag": "", "freq": "A", "unit": "PC_ENT"})
        lines.append(f"A,PC_ENT,{g}\t" + "\t".join(cells))
    write_text("eurostat_bulk.tsv", "\n".join(lines) + "\n")
    expect("eurostat_bulk.tsv", [["freq", "category"], ["unit", "category"], ["geo", "category"], ["Year", "time"],
                                 ["Value", "number"], ["flag", "category"]], n, spots,
           "Eurostat bulk TSV. Comma-joined codes in the first column, flags inside the cell like '12.3 p', ': ' for missing",
           aggregates=["EA20", "EU27_2020"], not_aggregates=["AT", "FR"])


def eurostat_tsv_size():
    rnd = random.Random(5)
    years = ["2021", "2023", "2024"]
    lines = ["freq,indic_is,unit,size_emp,geo\\TIME_PERIOD\t" + "\t".join(y + " " for y in years)]
    n, spots = 0, []
    for size in ["GE10", "GE250", "S10_49"]:
        for g in ["BE", "DK", "EU27_2020", "PL"]:
            cells = []
            for y in years:
                v = round(rnd.uniform(2, 60), 2)
                cells.append(f"{v} ")
                n += 1
                if g == "DK" and y == "2024":
                    spots.append({"size_emp": size, "geo": "DK", "Year": y, "Value": v, "indic_is": "E_AI_TANY"})
            lines.append(f"A,E_AI_TANY,PC_ENT,{size},{g}\t" + "\t".join(cells))
    write_text("eurostat_bulk_size.tsv", "\n".join(lines) + "\n")
    expect("eurostat_bulk_size.tsv", [["freq", "category"], ["indic_is", "category"], ["unit", "category"], ["size_emp", "category"],
                                      ["geo", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "Eurostat bulk TSV with five code dimensions and no flags", aggregates=["EU27_2020"])


def eurostat_sdmx_csv():
    rnd = random.Random(6)
    rows = [["DATAFLOW", "LAST UPDATE", "freq", "unit", "geo", "TIME_PERIOD", "OBS_VALUE", "OBS_FLAG", "CONF_STATUS"]]
    spots = []
    for g in ["AT", "BE", "DE", "EU27_2020"]:
        for y in range(2018, 2024):
            v = round(rnd.uniform(1, 9), 1)
            flag = "e" if (g == "DE" and y == 2020) else ""
            rows.append(["ESTAT:TEC00118(1.0)", "05/09/26 23:00:00", "A", "PC_GDP", g, y, v, flag, ""])
            if flag or (g == "AT" and y == 2018):
                spots.append({"geo": g, "TIME_PERIOD": str(y), "OBS_VALUE": v, "OBS_FLAG": flag})
    write_text("eurostat_sdmx.csv", csv_text(rows))
    expect("eurostat_sdmx.csv", [["DATAFLOW", "category"], ["LAST UPDATE", "time"], ["freq", "category"], ["unit", "category"],
                                 ["geo", "category"], ["TIME_PERIOD", "time"], ["OBS_VALUE", "number"], ["OBS_FLAG", "category"]],
           24, spots, "Eurostat SDMX-CSV, already long. CONF_STATUS is empty and dropped", aggregates=["EU27_2020"])


def oecd_stat():
    rnd = random.Random(7)
    head = ["LOCATION", "Country", "SUBJECT", "Subject", "MEASURE", "Measure", "FREQUENCY", "Frequency", "TIME", "Time",
            "Unit Code", "Unit", "PowerCode Code", "PowerCode", "Reference Period Code", "Reference Period", "Value",
            "Flag Codes", "Flags"]
    rows, spots = [head], []
    for code, name in [("AUT", "Austria"), ("BEL", "Belgium"), ("FRA", "France"), ("DEU", "Germany"), ("OECD", "OECD - Total")]:
        for y in range(2015, 2023):
            v = round(rnd.uniform(1, 4), 3)
            fc, fl = ("E", "Estimated value") if (code == "FRA" and y == 2022) else ("", "")
            rows.append([code, name, "GERD", "Gross domestic expenditure on R&D", "PC_GDP", "% of GDP", "A", "Annual", y, y,
                         "PC", "Percentage", 0, "Units", "", "", v, fc, fl])
            if fc or (code == "AUT" and y == 2015):
                spots.append({"LOCATION": code, "Country": name, "TIME": str(y), "Value": v, "Flag Codes": fc})
    write_text("oecd_stat_long.csv", csv_text(rows))
    cols = [["LOCATION", "category"], ["Country", "category"], ["SUBJECT", "category"], ["Subject", "category"],
            ["MEASURE", "category"], ["Measure", "category"], ["FREQUENCY", "category"], ["Frequency", "category"],
            ["TIME", "time"], ["Time", "time"], ["Unit Code", "category"], ["Unit", "category"], ["PowerCode Code", "number"],
            ["PowerCode", "category"], ["Value", "number"], ["Flag Codes", "category"], ["Flags", "category"]]
    expect("oecd_stat_long.csv", cols, 40, spots, "Legacy OECD.Stat CSV, code and label column pairs, mostly empty flag columns",
           aggregates=["OECD", "OECD - Total"], not_aggregates=["FRA"])


def oecd_explorer():
    rnd = random.Random(8)
    head = ["STRUCTURE", "STRUCTURE_ID", "STRUCTURE_NAME", "ACTION", "REF_AREA", "Reference area", "FREQ",
            "Frequency of observation", "MEASURE", "Measure", "UNIT_MEASURE", "Unit of measure", "TIME_PERIOD", "Time period",
            "OBS_VALUE", "Observation value", "OBS_STATUS", "Observation status", "UNIT_MULT", "Unit multiplier", "DECIMALS",
            "Decimals"]
    rows, spots = [head], []
    for code, name in [("BEL", "Belgium"), ("ITA", "Italy"), ("JPN", "Japan"), ("USA", "United States"), ("OECD", "OECD")]:
        for q in ["2023-Q1", "2023-Q2", "2023-Q3", "2023-Q4", "2024-Q1"]:
            v = round(rnd.uniform(50, 80), 1)
            st, stl = ("P", "Provisional value") if q == "2024-Q1" else ("A", "Normal value")
            rows.append(["DATAFLOW", "OECD.SDD.TPS:DSD_LFS@DF_IALFS_EMP_WAP_Q(1.0)", "Employment rate", "I", code, name, "Q",
                         "Quarterly", "EMP_WAP", "Employment rate", "PT_WAP", "Percentage of working age population", q, "", v,
                         "", st, stl, 0, "Units", 1, "One"])
            if code == "ITA" or (code == "OECD" and q == "2023-Q1"):
                spots.append({"REF_AREA": code, "TIME_PERIOD": q, "OBS_VALUE": v, "OBS_STATUS": st})
    write_text("oecd_explorer_long.csv", csv_text(rows))
    cols = [["STRUCTURE", "category"], ["STRUCTURE_ID", "category"], ["STRUCTURE_NAME", "category"], ["ACTION", "category"],
            ["REF_AREA", "category"], ["Reference area", "category"], ["FREQ", "category"], ["Frequency of observation", "category"],
            ["MEASURE", "category"], ["Measure", "category"], ["UNIT_MEASURE", "category"], ["Unit of measure", "category"],
            ["TIME_PERIOD", "time"], ["OBS_VALUE", "number"], ["OBS_STATUS", "category"], ["Observation status", "category"],
            ["UNIT_MULT", "number"], ["Unit multiplier", "category"], ["DECIMALS", "number"], ["Decimals", "category"]]
    expect("oecd_explorer_long.csv", cols, 25, spots,
           "OECD Data Explorer SDMX CSV with labels. Two label columns are empty and dropped", aggregates=["OECD"],
           not_aggregates=["ITA"])


def oecd_wide():
    rnd = random.Random(9)
    years = list(range(2016, 2024))
    rows = [["Dataset: Researchers per thousand employment, full-time equivalent"], ["Unit: Per 1 000 employed"],
            ["Country"] + [str(y) for y in years]]
    n, spots = 0, []
    for c in ["Australia", "Belgium", "Chile", "Korea", "Mexico", "OECD - Total", "European Union (27 countries from 01/02/2020)"]:
        r = [c]
        for y in years:
            if c == "Australia" and y % 2 == 1:
                r.append("..")
            else:
                v = round(rnd.uniform(1, 18), 1)
                r.append(v)
                n += 1
                if c in ("Korea", "Chile") and y == 2020:
                    spots.append({"Country": c, "Year": str(y), "Value": v})
        rows.append(r)
    rows.append(["Data extracted on 12 Oct 2026 09:40 UTC (GMT) from OECD.Stat"])
    write_text("oecd_wide.csv", csv_text(rows))
    expect("oecd_wide.csv", [["Country", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "OECD.Stat wide export with two title lines, '..' for missing, a trailing extraction line",
           aggregates=["OECD - Total", "European Union (27 countries from 01/02/2020)"], not_aggregates=["Korea"])


WB = [("Aruba", "ABW"), ("Africa Eastern and Southern", "AFE"), ("Austria", "AUT"), ("Belgium", "BEL"),
      ("Brazil", "BRA"), ("Central Europe and the Baltics", "CEB"), ("China", "CHN"), ("Euro area", "EMU"),
      ("East Asia & Pacific", "EAS"), ("European Union", "EUU"), ("France", "FRA"), ("Germany", "DEU"),
      ("High income", "HIC"), ("India", "IND"), ("Japan", "JPN"), ("Low & middle income", "LMY"), ("Middle income", "MIC"),
      ("OECD members", "OED"), ("Sub-Saharan Africa", "SSF"), ("United States", "USA"), ("World", "WLD"),
      ("Upper middle income", "UMC"), ("Small states", "SST"), ("Heavily indebted poor countries (HIPC)", "HPC")]
WB_AGG = ["Africa Eastern and Southern", "Central Europe and the Baltics", "Euro area", "East Asia & Pacific", "European Union",
          "High income", "Low & middle income", "Middle income", "OECD members", "Sub-Saharan Africa", "World",
          "Upper middle income", "Small states", "Heavily indebted poor countries (HIPC)"]


def wdi_rows(rnd):
    years = list(range(1960, 2025))
    head = ["Country Name", "Country Code", "Indicator Name", "Indicator Code"] + [str(y) for y in years]
    body, n, spots = [], 0, []
    for name, code in WB:
        r = [name, code, "Research and development expenditure (% of GDP)", "GB.XPD.RSDV.GD.ZS"]
        for y in years:
            if y < 1996 or (name == "Aruba") or (y == 2024 and name != "World"):
                r.append("")
            else:
                v = round(rnd.uniform(0.1, 4.5), 6)
                r.append(v)
                n += 1
                if name in ("Belgium", "World") and y == 2020:
                    spots.append({"Country Name": name, "Country Code": code, "Year": str(y), "Value": v})
        body.append(r)
    return head, body, n, spots


def wdi():
    head, body, n, spots = wdi_rows(random.Random(10))
    text = ('"Data Source","World Development Indicators",\n\n"Last Updated Date","2026-09-30",\n\n'
            + "".join(",".join(f'"{c}"' for c in r) + ",\n" for r in [head] + body))
    write_text("worldbank_wdi.csv", text)
    expect("worldbank_wdi.csv", [["Country Name", "category"], ["Country Code", "category"], ["Indicator Name", "category"],
                                 ["Indicator Code", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "World Bank WDI CSV. Four metadata lines with blanks, 65 year columns, region and income aggregates, trailing comma",
           aggregates=WB_AGG, not_aggregates=["Belgium", "United States", "China"])


def wdi_xlsx():
    head, body, n, spots = wdi_rows(random.Random(11))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Data Source", "World Development Indicators"])
    ws.append([])
    ws.append(["Last Updated Date", dt.datetime(2026, 9, 30)])
    ws.append([])
    ws.append(head)
    for r in body:
        ws.append([None if c == "" else c for c in r])
    m = wb.create_sheet("Metadata - Countries")
    m.append(["Country Code", "Region", "IncomeGroup", "SpecialNotes", "TableName"])
    for name, code in WB:
        m.append([code, "" if name in WB_AGG else "Europe & Central Asia", "" if name in WB_AGG else "High income", "", name])
    m2 = wb.create_sheet("Metadata - Indicators")
    m2.append(["INDICATOR_CODE", "INDICATOR_NAME", "SOURCE_NOTE", "SOURCE_ORGANIZATION"])
    m2.append(["GB.XPD.RSDV.GD.ZS", "Research and development expenditure (% of GDP)", "Gross domestic expenditures ...",
               "UNESCO Institute for Statistics"])
    save_wb(wb, "worldbank_wdi.xlsx")
    expect("worldbank_wdi.xlsx", [["Country Name", "category"], ["Country Code", "category"], ["Indicator Name", "category"],
                                  ["Indicator Code", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "World Bank Excel download. Data sheet first, two metadata sheets", aggregates=WB_AGG, not_aggregates=["Belgium"],
           sheet="Data")


def ilostat():
    rnd = random.Random(12)
    head = ["ref_area.label", "source.label", "indicator.label", "sex.label", "classif1.label", "time", "obs_value",
            "obs_status.label", "note_classif.label", "note_indicator.label", "note_source.label"]
    rows, spots = [head], []
    for area in ["Belgium", "France", "Germany", "Italy"]:
        for sex in ["Sex: Total", "Sex: Male", "Sex: Female"]:
            for age in ["Age (Youth, adults): 15+", "Age (Youth, adults): 15-24"]:
                for y in [2022, 2023, 2024]:
                    v = round(rnd.uniform(2, 30), 3)
                    st = "Break in series" if (area == "Germany" and y == 2024) else ""
                    rows.append([area, "LFS - EU Labour Force Survey", "Unemployment rate by sex and age (%)", sex, age, y, v, st,
                                 "", "Repository: ILO-STATISTICS - Micro data processing", "Break in series: Methodology revised" if st else ""])
                    if area == "France" and sex == "Sex: Female" and y == 2023:
                        spots.append({"ref_area.label": area, "sex.label": sex, "classif1.label": age, "time": "2023", "obs_value": v})
    write_text("ilostat_bulk.csv", csv_text(rows))
    expect("ilostat_bulk.csv", [["ref_area.label", "category"], ["source.label", "category"], ["indicator.label", "category"],
                                ["sex.label", "category"], ["classif1.label", "category"], ["time", "time"],
                                ["obs_value", "number"], ["obs_status.label", "category"], ["note_indicator.label", "category"],
                                ["note_source.label", "category"]], 72, spots,
           "ILOSTAT bulk CSV, long, with label columns and a mostly empty status column", aggregates=["Sex: Total"],
           not_aggregates=["Sex: Male", "Belgium"])


def statbel():
    rnd = random.Random(13)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tableau 4"
    ws.append(["Tableau 4. Entreprises utilisant l'intelligence artificielle, selon la région et la taille, 2025"])
    ws.append([])
    ws.append(["Région / Province", "Nombre d'entreprises", None, None, "Part utilisant l'IA (%)", None, None])
    ws.append([None, "10-49", "50-249", "250+", "10-49", "50-249", "250+"])
    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:D3")
    ws.merge_cells("E3:G3")
    prov = {"Région flamande": ["Anvers", "Limbourg", "Flandre orientale", "Brabant flamand", "Flandre occidentale"],
            "Région wallonne": ["Brabant wallon", "Hainaut", "Liège", "Luxembourg", "Namur"]}
    leaf = {}
    for p in prov["Région flamande"] + prov["Région wallonne"] + ["Région de Bruxelles-Capitale"]:
        leaf[p] = [rnd.randint(300, 4000), rnd.randint(40, 600), rnd.randint(5, 120)]
    def tot(names):
        return [sum(leaf[x][k] for x in names) for k in range(3)]
    order = [("Belgique", tot(list(leaf))), ("Région de Bruxelles-Capitale¹", leaf["Région de Bruxelles-Capitale"])]
    for reg in ["Région flamande", "Région wallonne"]:
        order.append((reg, tot(prov[reg])))
        order += [(p, leaf[p]) for p in prov[reg]]
    expected_rows = []
    for name, counts in order:
        shares = [round(rnd.uniform(5, 50), 1) for _ in range(3)]
        cells_sh = list(shares)
        if name == "Hainaut":
            cells_sh[2] = f"{str(shares[2]).replace('.', ',')}²"
        if name == "Namur":
            cells_sh[2] = "c"
            shares[2] = None
        ws.append([name] + counts + cells_sh)
        expected_rows.append((name.rstrip("¹"), counts, shares))
    ws.append([])
    ws.append(["¹ Y compris les entreprises dont le siège est à Bruxelles mais l'activité en Flandre."])
    ws.append(["² Estimation peu fiable."])
    ws.append(["c : confidentiel."])
    ws.append(["Source : Statbel (Enquête TIC 2025), données fictives"])
    save_wb(wb, "statbel_merged_headers.xlsx")
    names = ["Région / Province", "Nombre d'entreprises 10-49", "Nombre d'entreprises 50-249", "Nombre d'entreprises 250+",
             "Part utilisant l'IA (%) 10-49", "Part utilisant l'IA (%) 50-249", "Part utilisant l'IA (%) 250+"]
    spots = []
    for name, counts, shares in expected_rows:
        if name in ("Région de Bruxelles-Capitale", "Hainaut", "Namur", "Belgique"):
            spots.append(dict(zip(names, [name] + counts + shares)))
    expect("statbel_merged_headers.xlsx", [[names[0], "category"]] + [[n, "number"] for n in names[1:]], len(expected_rows), spots,
           "Statbel-style table. Title, two-row header with merged cells, footnote marks in names and values, subtotals",
           aggregates=["Belgique", "Région flamande", "Région wallonne"],
           not_aggregates=["Région de Bruxelles-Capitale", "Hainaut", "Anvers"])


def genesis():
    rnd = random.Random(14)
    lines = ["GENESIS-Tabelle: 52911-0004;;;", "Unternehmen mit Nutzung von Künstlicher Intelligenz: Deutschland,;;;",
             "Jahre, Wirtschaftszweige (WZ2008);;;", "Erhebung zur Nutzung von IKT in Unternehmen;;;", ";;;",
             ";Anteil der Unternehmen;Anteil der Unternehmen;Anteil der Unternehmen", ";Prozent;Prozent;Prozent",
             "Wirtschaftszweige (WZ2008);2021;2023;2024"]
    sectors = ["WZ08-C Verarbeitendes Gewerbe", "WZ08-D Energieversorgung", "WZ08-F Baugewerbe", "WZ08-G Handel",
               "WZ08-J Information und Kommunikation", "WZ08-M Freiberufliche Dienstleistungen", "Insgesamt"]
    n, spots = 0, []
    for s in sectors:
        cells = []
        for k, y in enumerate(["2021", "2023", "2024"]):
            v = round(rnd.uniform(3, 40), 1)
            if s == "WZ08-D Energieversorgung" and k == 0:
                cells.append(".")
            elif s == "WZ08-F Baugewerbe" and k == 1:
                cells.append("-")
            elif s == "WZ08-G Handel" and k == 2:
                cells.append("x")
            else:
                cells.append(str(v).replace(".", ","))
                n += 1
                if k == 1 and s in ("WZ08-J Information und Kommunikation", "Insgesamt"):
                    spots.append({"Wirtschaftszweige (WZ2008)": s, "Year": y, "Value": v})
        lines.append(s + ";" + ";".join(cells))
    lines += ["__________", "Zeichenerklärung: . = Zahlenwert unbekannt oder geheim zu halten, - = nichts vorhanden, x = Tabellenfach gesperrt",
              "Stand: 15.09.2026 / 11:04:52", "(C)opyright Statistisches Bundesamt (Destatis), 2026 | Stand: 15.09.2026"]
    write_text("destatis_genesis.csv", "\n".join(lines) + "\n", encoding="cp1252")
    expect("destatis_genesis.csv", [["Wirtschaftszweige (WZ2008)", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "Destatis GENESIS CSV in cp1252. Semicolons, decimal comma, title lines, three header rows with constant measure and unit rows",
           aggregates=["Insgesamt"], not_aggregates=["WZ08-G Handel"])


FR_MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
DE_MONTHS = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]


def french_export():
    rnd = random.Random(15)
    lines = ['"Offres d\'emploi en ligne mentionnant l\'intelligence artificielle, par région";;;',
             '"Champ : France métropolitaine, données fictives";;;', ";;;", "Mois;Région;Offres d'emploi;Part IA (%)"]
    n, spots = 0, []
    for m in range(12):
        for reg in ["Île-de-France", "Auvergne-Rhône-Alpes", "Occitanie", "Hauts-de-France", "Bretagne"]:
            offers = rnd.randint(1200, 98000)
            share = round(rnd.uniform(0.5, 6), 1)
            off_txt = f"{offers:,}".replace(",", " ")
            sh_txt = str(share).replace(".", ",")
            if reg == "Bretagne" and m == 3:
                sh_txt, share = "n.d.", None
            lines.append(f"{FR_MONTHS[m]} 2024;{reg};{off_txt};{sh_txt}")
            n += 1
            if (reg == "Île-de-France" and m == 1) or (reg == "Bretagne" and m == 3):
                spots.append({"Mois": f"2024-{m + 1:02d}", "Région": reg, "Offres d'emploi": offers, "Part IA (%)": share})
    write_text("insee_style_semicolon.csv", "\n".join(lines) + "\n")
    expect("insee_style_semicolon.csv", [["Mois", "time"], ["Région", "category"], ["Offres d'emploi", "number"],
                                         ["Part IA (%)", "number"]], n, spots,
           "French export. Semicolons, decimal comma, thousands with narrow no-break spaces, month names in French, n.d.")


def german_export():
    rnd = random.Random(16)
    lines = ["Monat;Bundesland;Stellenanzeigen;KI-Anteil in %"]
    n, spots = 0, []
    for m in range(12):
        for land in ["Bayern", "Berlin", "Hessen", "Sachsen"]:
            ads = rnd.randint(1500, 250000)
            sh = round(rnd.uniform(0.5, 7), 2)
            ads_txt = f"{ads:,}".replace(",", ".")
            lines.append(f"{DE_MONTHS[m]} 2025;{land};{ads_txt};{str(sh).replace('.', ',')}")
            n += 1
            if land == "Berlin" and m in (2, 10):
                spots.append({"Monat": f"2025-{m + 1:02d}", "Bundesland": land, "Stellenanzeigen": ads, "KI-Anteil in %": sh})
    write_text("german_semicolon_thousands.csv", "\n".join(lines) + "\r\n", newline="\r\n")
    expect("german_semicolon_thousands.csv", [["Monat", "time"], ["Bundesland", "category"], ["Stellenanzeigen", "number"],
                                              ["KI-Anteil in %", "number"]], n, spots,
           "German export. Semicolons, thousands with dots, decimal comma, German month names, Windows line endings")


def excel_report():
    rnd = random.Random(17)
    wb = openpyxl.Workbook()
    rd = wb.active
    rd.title = "Read me"
    for t in ["About this workbook", "", "This workbook holds invented figures for testing.",
              "Sheet Data has one row per region and quarter end.", "Contact: data team"]:
        rd.append([t])
    ws = wb.create_sheet("Data")
    ws.append(["AI adoption in EU regions"])
    ws.append(["Percent of enterprises, by quarter end"])
    ws.append([])
    ws.append(["Region", None, "Date", "Share", "Firms"])
    spots, n = [], 0
    for reg in ["Brussels", "Flanders", "Wallonia", "Île-de-France", "Bavaria"]:
        for q, d in enumerate([dt.datetime(2025, 3, 31), dt.datetime(2025, 6, 30), dt.datetime(2025, 9, 30)]):
            sh = round(rnd.uniform(5, 35), 1)
            firms = rnd.randint(100, 9000)
            ws.append([reg, None, d, f"{sh}%", firms])
            ws.cell(ws.max_row, 3).number_format = "dd/mm/yyyy"
            n += 1
            if reg == "Wallonia":
                spots.append({"Region": reg, "Date": d.date().isoformat(), "Share": sh, "Firms": firms})
    ws.append(["Total", None, None, "", 99999])
    n += 1
    ws.append([])
    ws.append(["Notes: Share is the percent of enterprises with 10 or more persons employed."])
    save_wb(wb, "excel_report_notes_sheet.xlsx")
    expect("excel_report_notes_sheet.xlsx", [["Region", "category"], ["Date", "time"], ["Share", "number"], ["Firms", "number"]], n,
           spots + [{"Region": "Total", "Firms": 99999, "Share": None}],
           "Excel report. A read-me sheet comes first, title block, an empty spacer column, dates as Excel dates, percent as text, a totals row",
           aggregates=["Total"], not_aggregates=["Wallonia"], sheet="Data")


def excel_serials():
    rnd = random.Random(18)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "export"
    ws.append(["date", "series", "index_value"])
    spots = []
    base = dt.date(1899, 12, 30)
    for k in range(24):
        d = dt.date(2024, 1, 1) + dt.timedelta(days=7 * k)
        serial = (d - base).days
        for s in ["AI stocks", "Market"]:
            v = round(rnd.uniform(90, 140), 2)
            ws.append([serial, s, v])
            if k in (0, 10):
                spots.append({"date": d.isoformat(), "series": s, "index_value": v})
    save_wb(wb, "excel_serial_dates.xlsx")
    expect("excel_serial_dates.xlsx", [["date", "time"], ["series", "category"], ["index_value", "number"]], 48, spots,
           "Dates stored as raw Excel serial numbers in a column named date")


def crosstab_csv():
    rnd = random.Random(19)
    rows = [["Employment rate of people aged 20-64 by sex (%)"], ["Country", "2022", "", "2023", "", "2024", ""],
            ["", "Men", "Women", "Men", "Women", "Men", "Women"]]
    n, spots = 0, []
    for c in EU[:8] + ["EU27"]:
        r = [c]
        for y in ["2022", "2023", "2024"]:
            for s in ["Men", "Women"]:
                if c == "Greece" and y == "2022" and s == "Women":
                    r.append(":")
                    continue
                v = round(rnd.uniform(55, 85), 1)
                r.append(v)
                n += 1
                if c == "Germany" and y == "2023":
                    spots.append({"Country": c, "Year": y, "Sex": s, "Value": v})
        rows.append(r)
    write_text("crosstab_country_year_sex.csv", csv_text(rows))
    expect("crosstab_country_year_sex.csv", [["Country", "category"], ["Year", "time"], ["Sex", "category"], ["Value", "number"]],
           n, spots, "Crosstab with two header rows. Years span two columns each, sex below", aggregates=["EU27"],
           not_aggregates=["Germany"])


def crosstab_xlsx():
    rnd = random.Random(20)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table"
    ws.append(["Tertiary graduates in ICT by sex"])
    ws.append([])
    ws.append(["Country", 2021, None, None, 2022, None, None])
    ws.append([None, "Males", "Females", "Total", "Males", "Females", "Total"])
    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:D3")
    ws.merge_cells("E3:G3")
    n, spots = 0, []
    for c in EU[8:16]:
        r = [c]
        for y in ["2021", "2022"]:
            m, f = rnd.randint(500, 9000), rnd.randint(100, 3000)
            r += [m, f, m + f]
            n += 3
            if c == "Italy":
                spots += [{"Country": c, "Year": y, "Sex": "Females", "Value": f}, {"Country": c, "Year": y, "Sex": "Total", "Value": m + f}]
        ws.append(r)
    save_wb(wb, "crosstab_merged.xlsx")
    expect("crosstab_merged.xlsx", [["Country", "category"], ["Year", "time"], ["Sex", "category"], ["Value", "number"]], n, spots,
           "Crosstab in Excel with merged year cells over three sex columns", aggregates=["Total"], not_aggregates=["Italy"])


def clean_csv():
    rnd = random.Random(21)
    rows = [["country", "year", "gdp_growth", "inflation"]]
    spots = []
    for c in EU[:10]:
        for y in [2022, 2023, 2024]:
            g, i = round(rnd.uniform(-2, 6), 1), round(rnd.uniform(0, 12), 1)
            rows.append([c, y, g, i])
            if c == "Spain":
                spots.append({"country": c, "year": str(y), "gdp_growth": g, "inflation": i})
    write_text("clean_long.csv", csv_text(rows))
    expect("clean_long.csv", [["country", "category"], ["year", "time"], ["gdp_growth", "number"], ["inflation", "number"]], 30,
           spots, "A clean long CSV with the header on line one", not_aggregates=["Spain"])


def clean_xlsx():
    rnd = random.Random(22)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Company", "Country", "Founded", "Employees", "Funding (EUR m)"])
    spots = []
    for k in range(40):
        row = [f"Startup {k + 1:02d}", rnd.choice(EU[:6]), rnd.randint(2008, 2024), rnd.randint(3, 900), round(rnd.uniform(0.2, 300), 1)]
        ws.append(row)
        if k in (0, 39):
            spots.append(dict(zip(["Company", "Country", "Founded", "Employees", "Funding (EUR m)"], [row[0], row[1], str(row[2]), row[3], row[4]])))
    save_wb(wb, "clean_simple.xlsx")
    expect("clean_simple.xlsx", [["Company", "category"], ["Country", "category"], ["Founded", "time"], ["Employees", "number"],
                                 ["Funding (EUR m)", "number"]], 40, spots, "A clean Excel sheet with the header in row 1")


def json_records():
    rnd = random.Random(23)
    recs, spots = [], []
    for c in EU[:6]:
        for y in [2023, 2024]:
            v = None if (c == "Denmark" and y == 2024) else round(rnd.uniform(1, 30), 2)
            recs.append({"country": c, "year": y, "ai_share": v, "source": "mock"})
            if c in ("Denmark", "Bulgaria"):
                spots.append({"country": c, "year": str(y), "ai_share": v})
    write_text("records.json", json.dumps(recs, indent=1))
    expect("records.json", [["country", "category"], ["year", "time"], ["ai_share", "number"], ["source", "category"]], 12, spots,
           "JSON array of records, one null")


def json_stat():
    rnd = random.Random(24)
    geos = {"BE": "Belgium", "DE": "Germany", "FR": "France", "EU27_2020": "European Union - 27 countries (from 2020)"}
    years = ["2022", "2023", "2024"]
    vals, spots = [], []
    for gk, gl in geos.items():
        for y in years:
            v = None if (gk == "FR" and y == "2022") else round(rnd.uniform(5, 25), 1)
            vals.append(v)
            if gk == "DE" or v is None:
                spots.append({"Geopolitical entity": gl, "Time": y, "Value": v})
    doc = {"version": "2.0", "class": "dataset", "label": "Enterprises using AI (mock)", "source": "Eurostat (mock)",
           "id": ["unit", "geo", "time"], "size": [1, 4, 3],
           "dimension": {"unit": {"label": "Unit of measure", "category": {"index": {"PC_ENT": 0}, "label": {"PC_ENT": "Percentage of enterprises"}}},
                         "geo": {"label": "Geopolitical entity", "category": {"index": {k: i for i, k in enumerate(geos)}, "label": geos}},
                         "time": {"label": "Time", "category": {"index": {y: i for i, y in enumerate(years)}, "label": {y: y for y in years}}}},
           "value": vals}
    write_text("jsonstat_dataset.json", json.dumps(doc, ensure_ascii=False))
    expect("jsonstat_dataset.json", [["Unit of measure", "category"], ["Geopolitical entity", "category"], ["Time", "time"],
                                     ["Value", "number"]], 12, spots, "JSON-stat 2.0 dataset as the Eurostat API returns it",
           aggregates=["European Union - 27 countries (from 2020)"])


def bom_csv():
    rows = [["Région", "Année", "Valeur"]]
    spots = []
    rnd = random.Random(25)
    for reg in ["Bruxelles", "Région wallonne", "Vlaams Gewest"]:
        for y in [2023, 2024]:
            v = round(rnd.uniform(1, 9), 2)
            rows.append([reg, y, v])
            spots.append({"Région": reg, "Année": str(y), "Valeur": v})
    write_text("utf8_bom.csv", "﻿" + csv_text(rows))
    expect("utf8_bom.csv", [["Région", "category"], ["Année", "time"], ["Valeur", "number"]], 6, spots[:2],
           "UTF-8 with a byte order mark before the first header")


def latin1_csv():
    rows = [["Département;Année;Taux de chômage"]]
    rnd = random.Random(26)
    spots = []
    lines = ["Département;Année;Taux de chômage"]
    for d in ["Côte-d'Or", "Hérault", "Bouches-du-Rhône", "Pyrénées-Atlantiques"]:
        for y in [2023, 2024]:
            v = round(rnd.uniform(5, 12), 1)
            lines.append(f"{d};{y};{str(v).replace('.', ',')}")
            if d in ("Côte-d'Or", "Hérault"):
                spots.append({"Département": d, "Année": str(y), "Taux de chômage": v})
    write_text("latin1_semicolon.csv", "\n".join(lines) + "\n", encoding="latin-1")
    expect("latin1_semicolon.csv", [["Département", "category"], ["Année", "time"], ["Taux de chômage", "number"]], 8, spots,
           "Latin-1 text with accents, semicolons and decimal commas")


def crlf_csv():
    rows = [["Organisation", "Country", "Projects", "EU contribution (EUR)"]]
    rnd = random.Random(27)
    spots = []
    for k in range(15):
        org = f"Institute of Applied Research, Unit {k + 1}"
        c = EU[k % 9]
        p, e = rnd.randint(1, 40), round(rnd.uniform(1e5, 9e7), 2)
        rows.append([org, c, p, f"{e:,.2f}"])
        if k in (0, 14):
            spots.append({"Organisation": org, "Country": c, "Projects": p, "EU contribution (EUR)": e})
    write_text("windows_crlf_quoted.csv", csv_text(rows), newline="\r\n")
    expect("windows_crlf_quoted.csv", [["Organisation", "category"], ["Country", "category"], ["Projects", "number"],
                                       ["EU contribution (EUR)", "number"]], 15, spots,
           "CORDIS-style CSV with Windows line endings, quoted names holding commas, thousands commas")


def wide60():
    rnd = random.Random(28)
    head = ["country"] + [f"ind_{k:02d}" for k in range(1, 60)]
    rows, spots = [head], []
    for c in EU[:20]:
        r = [c] + [round(rnd.uniform(0, 100), 2) for _ in range(59)]
        rows.append(r)
        if c == "Germany":
            spots.append({"country": c, "ind_01": r[1], "ind_33": r[33], "ind_59": r[59]})
    write_text("very_wide_60_columns.csv", csv_text(rows))
    expect("very_wide_60_columns.csv", [["country", "category"]] + [[h, "number"] for h in head[1:]], 20, spots,
           "60 columns, one per indicator, not years")


def long50k():
    rnd = random.Random(29)
    rows = [["nuts2", "isco08", "year", "postings", "ai_postings"]]
    spots = []
    regs = [f"{c}{k}{j}" for c in ISO2[:10] for k in range(1, 6) for j in range(1, 6)]
    iscos = [f"OC{k}" for k in range(1, 11)]
    for reg in regs:
        for isco in iscos:
            for y in range(2005, 2025):
                p = rnd.randint(10, 20000)
                a = rnd.randint(0, p // 10)
                rows.append([reg, isco, y, p, a])
                if reg == "DE34" and isco == "OC2" and y > 2020:
                    spots.append({"nuts2": reg, "isco08": isco, "year": str(y), "postings": p, "ai_postings": a})
    write_text("long_50k_rows.csv", csv_text(rows))
    expect("long_50k_rows.csv", [["nuts2", "category"], ["isco08", "category"], ["year", "time"], ["postings", "number"],
                                 ["ai_postings", "number"]], len(rows) - 1, spots, "50,000 rows of online job ads, long format")


def repeated_header():
    rnd = random.Random(30)
    lines, n, spots = [], 0, []
    for sex in ["Total", "Females", "Males"]:
        lines.append(f"Sex: {sex}")
        lines.append("geo,2021,2022,2023")
        for g in ["BE", "DE", "FR", "NL"]:
            vals = [round(rnd.uniform(60, 80), 1) for _ in range(3)]
            lines.append(g + "," + ",".join(map(str, vals)))
            n += 3
            if g == "FR":
                spots.append({"Sex": sex, "geo": g, "Year": "2022", "Value": vals[1]})
    write_text("repeated_header_blocks.csv", "\n".join(lines) + "\n")
    expect("repeated_header_blocks.csv", [["Sex", "category"], ["geo", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "Three tables in one CSV with no blank lines. Each starts with a 'Sex: ...' line and repeats the header",
           aggregates=["Total"], not_aggregates=["FR"])


def stacked_csv():
    rnd = random.Random(31)
    rows, n, spots = [["Enterprises with e-commerce sales"], [], ["Last update", "01/10/2026"], []], 0, []
    for unit in ["Percentage of enterprises", "Percentage of turnover"]:
        rows += [["Unit of measure", unit], [], ["TIME", "2022", "2023", "2024"], ["GEO (Labels)", "", "", ""]]
        for g in ["European Union - 27 countries (from 2020)", "Belgium", "Czechia", "Ireland", "Portugal"]:
            r = [g]
            for y in ["2022", "2023", "2024"]:
                v = round(rnd.uniform(5, 40), 1)
                r.append(v)
                n += 1
                if g == "Ireland":
                    spots.append({"Unit of measure": unit, "GEO": g, "Year": y, "Value": v})
            rows.append(r)
        rows += [[], []]
    rows += [["Special value"], [":", "not available"]]
    write_text("eurostat_stacked_units.csv", csv_text(rows))
    expect("eurostat_stacked_units.csv", [["Unit of measure", "category"], ["GEO", "category"], ["Year", "time"], ["Value", "number"]],
           n, spots, "Eurostat-style CSV with two tables, one per unit of measure, split by blank rows",
           aggregates=["European Union - 27 countries (from 2020)"], not_aggregates=["Ireland"])


def text_only():
    rows = [["country", "capital", "region", "eu_member"]]
    for c, cap, reg, m in [("Belgium", "Brussels", "Western Europe", "yes"), ("Norway", "Oslo", "Northern Europe", "no"),
                           ("Poland", "Warsaw", "Central Europe", "yes"), ("Serbia", "Belgrade", "Southern Europe", "no"),
                           ("Portugal", "Lisbon", "Southern Europe", "yes")]:
        rows.append([c, cap, reg, m])
    write_text("text_only.csv", csv_text(rows))
    expect("text_only.csv", [["country", "category"], ["capital", "category"], ["region", "category"], ["eu_member", "category"]], 5,
           [{"country": "Norway", "capital": "Oslo", "eu_member": "no"}], "No numbers at all")


def quarters_wide():
    rnd = random.Random(32)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quarterly"
    qs = [f"{y}Q{q}" for y in (2023, 2024) for q in range(1, 5)]
    ws.append(["Online job ads requiring AI skills, by region (thousands)"])
    ws.append(["Region"] + qs)
    n, spots = 0, []
    for reg in ["Île-de-France", "Lombardia", "Oberbayern", "Comunidad de Madrid", "Noord-Holland", "Mazowieckie"]:
        r = [reg]
        for q in qs:
            v = round(rnd.uniform(0.5, 25), 2)
            r.append(v)
            n += 1
            if reg == "Lombardia" and q in ("2023Q2", "2024Q4"):
                spots.append({"Region": reg, "Period": q, "Value": v})
        ws.append(r)
    ws.append(["Source: Lightcast, mock figures"])
    save_wb(wb, "quarters_wide.xlsx")
    expect("quarters_wide.xlsx", [["Region", "category"], ["Period", "time"], ["Value", "number"]], n, spots,
           "Quarterly periods as columns like 2023Q1, title and source lines")


def lightcast():
    rnd = random.Random(33)
    rows = [["NUTS2", "Region", "Year", "Postings", "AI postings", "AI share (%)"]]
    spots, tp, ta = [], 0, 0
    for code, name in [("BE10", "Région de Bruxelles-Capitale"), ("BE21", "Prov. Antwerpen"), ("BE22", "Prov. Limburg (BE)"),
                       ("BE23", "Prov. Oost-Vlaanderen"), ("BE24", "Prov. Vlaams-Brabant"), ("BE31", "Prov. Brabant wallon")]:
        p = rnd.randint(20000, 400000)
        a = rnd.randint(200, 9000)
        tp, ta = tp + p, ta + a
        rows.append([code, name, 2025, f"{p:,}", f"{a:,}", round(100 * a / p, 2)])
        if code == "BE22":
            spots.append({"NUTS2": code, "Region": name, "Postings": p, "AI postings": a})
    rows.append(["BE", "Total", 2025, f"{tp:,}", f"{ta:,}", round(100 * ta / tp, 2)])
    write_text("lightcast_regions.csv", csv_text(rows))
    expect("lightcast_regions.csv", [["NUTS2", "category"], ["Region", "category"], ["Year", "time"], ["Postings", "number"],
                                     ["AI postings", "number"], ["AI share (%)", "number"]], 7,
           spots + [{"NUTS2": "BE", "Region": "Total", "Postings": tp}],
           "Lightcast-style regional extract with thousands commas and a total row", aggregates=["BE"], not_aggregates=["BE22"])


def nace_sum():
    rnd = random.Random(34)
    secs = [("C", "Manufacturing"), ("F", "Construction"), ("G", "Wholesale and retail trade"), ("H", "Transportation and storage"),
            ("J", "Information and communication"), ("M", "Professional, scientific and technical activities")]
    rows = [["nace_r2", "sector", "2022", "2023", "2024"]]
    data = {c: [rnd.randint(1000, 90000) for _ in range(3)] for c, _ in secs}
    tot = [sum(data[c][k] for c, _ in secs) for k in range(3)]
    rows.append(["B-S", "Industry, construction and services", *tot])
    for c, s in secs:
        rows.append([c, s, *data[c]])
    write_text("nace_sum_check.csv", csv_text(rows))
    expect("nace_sum_check.csv", [["nace_r2", "category"], ["sector", "category"], ["Year", "time"], ["Value", "number"]], 21,
           [{"nace_r2": "B-S", "Year": "2023", "Value": tot[1]}, {"nace_r2": "J", "Year": "2024", "Value": data["J"][2]}],
           "Enterprises by NACE section with a B-S row that only a sum check reveals", aggregates=["B-S"],
           not_aggregates=["J", "C"])


def pandas_index():
    rnd = random.Random(35)
    rows = [["", "country", "score", "rank"]]
    spots = []
    for k, c in enumerate(EU[:12]):
        s = round(rnd.uniform(20, 90), 3)
        rows.append([k, c, s, k + 1])
        if k == 5:
            spots.append({"country": c, "score": s, "rank": 6})
    write_text("pandas_index_column.csv", csv_text(rows))
    expect("pandas_index_column.csv", [["country", "category"], ["score", "number"], ["rank", "number"]], 12, spots,
           "pandas to_csv with the index written as an unnamed first column")


def merged_rows():
    rnd = random.Random(36)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Horizon Europe budget by cluster and destination (EUR million)"])
    ws.append([])
    ws.append(["Cluster", "Destination", "2025", "2026", "2027"])
    spots, n, r0 = [], 0, 4
    for cl, dests in [("Digital, Industry and Space", ["AI and robotics", "Data and computing", "Space"]),
                      ("Health", ["Staying healthy", "Health care"]), ("Climate, Energy and Mobility", ["Energy supply", "Mobility", "Batteries"])]:
        start = r0
        for k, d in enumerate(dests):
            vals = [round(rnd.uniform(50, 900), 1) for _ in range(3)]
            ws.append([cl if k == 0 else None, d] + vals)
            n += 3
            if d in ("Data and computing", "Batteries"):
                spots.append({"Cluster": cl, "Destination": d, "Year": "2026", "Value": vals[1]})
            r0 += 1
        ws.merge_cells(start_row=start, start_column=1, end_row=r0 - 1, end_column=1)
    save_wb(wb, "merged_category_rows.xlsx")
    expect("merged_category_rows.xlsx", [["Cluster", "category"], ["Destination", "category"], ["Year", "time"], ["Value", "number"]],
           n, spots, "A category merged down several rows. Every row needs the cluster name")


def pipe_csv():
    rnd = random.Random(37)
    lines = ["patent_office|applicant_country|ipc_class|year|applications"]
    spots = []
    for off in ["EPO", "USPTO"]:
        for c in ["DE", "FR", "CN", "US"]:
            for y in [2021, 2022]:
                v = rnd.randint(10, 5000)
                lines.append(f"{off}|{c}|G06N|{y}|{v}")
                if off == "EPO" and c == "CN":
                    spots.append({"patent_office": off, "applicant_country": c, "year": str(y), "applications": v})
    write_text("pipe_delimited.txt", "\n".join(lines) + "\n")
    expect("pipe_delimited.txt", [["patent_office", "category"], ["applicant_country", "category"], ["ipc_class", "category"],
                                  ["year", "time"], ["applications", "number"]], 16, spots, "Pipe-delimited text export")


def footnote_header_csv():
    rnd = random.Random(38)
    rows = [["Table 2.1 Public R&D spending on AI, EUR million"], [],
            ["Member State", "2022", "2023", "2024*"]]
    n, spots = 0, []
    for c in EU[:7] + ["EU-27"]:
        vals = [round(rnd.uniform(10, 900), 1) for _ in range(3)]
        cells = list(vals)
        if c == "Czechia":
            cells[2] = f"{vals[2]}e"
        rows.append([c] + cells)
        n += 3
        if c in ("Czechia", "EU-27"):
            spots.append({"Member State": c, "Year": "2024", "Value": vals[2], "flag": "e" if c == "Czechia" else ""})
    rows += [[], ["* Provisional."], ["e Estimated."], ["Source: national budgets, mock figures"]]
    write_text("report_table_footnotes.csv", csv_text(rows))
    expect("report_table_footnotes.csv", [["Member State", "category"], ["Year", "time"], ["Value", "number"], ["flag", "category"]], n, spots,
           "Report table pasted to CSV. A year header with a footnote star, an estimate flag glued to a value",
           aggregates=["EU-27"], not_aggregates=["Czechia"])


def side_by_side():
    rnd = random.Random(39)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["Men", None, None, None, None, "Women"])
    ws.append(["Country", 2022, 2023, 2024, None, "Country", 2022, 2023, 2024])
    for c in EU[:9]:
        ws.append([c] + [round(rnd.uniform(60, 85), 1) for _ in range(3)] + [None, c] + [round(rnd.uniform(50, 80), 1) for _ in range(3)])
    save_wb(wb, "side_by_side_tables.xlsx")
    expect("side_by_side_tables.xlsx", [["Sex", "category"], ["Country", "category"], ["Year", "time"], ["Value", "number"]], 54, [],
           "Two tables next to each other on one sheet. Rules cannot split them by columns", needs_ai=True)


def no_header():
    rnd = random.Random(40)
    rows = [[c, y, round(rnd.uniform(1, 9), 2)] for c in EU[:6] for y in (2023, 2024)]
    write_text("no_header_row.csv", csv_text(rows))
    expect("no_header_row.csv", [["country", "category"], ["year", "time"], ["value", "number"]], 12, [],
           "Data with no header row at all. Only a reader of the content can name the columns", needs_ai=True)


def group_rows():
    rnd = random.Random(41)
    groups = [("Industrie", ["Industrie manufacturière", "Énergie, eau, déchets"]), ("Construction", ["Bâtiment", "Travaux publics"]),
              ("Services", ["Commerce", "Information et communication", "Services aux entreprises"])]
    rows = [["Tableau 2 - Emploi salarié par secteur (milliers)"], ["Secteur", "2022", "2023", "2024"]]
    n, spots = 0, []
    for g, subs in groups:
        rows.append([g])
        for sub in subs:
            vals = [round(rnd.uniform(100, 3000), 1) for _ in range(3)]
            rows.append(["  " + sub] + vals)
            n += 3
            if sub in ("Bâtiment", "Commerce"):
                spots.append({"Group": g, "Secteur": sub, "Year": "2023", "Value": vals[1]})
    write_text("section_heading_rows.csv", csv_text(rows))
    expect("section_heading_rows.csv", [["Group", "category"], ["Secteur", "category"], ["Year", "time"], ["Value", "number"]], n, spots,
           "A national table where section headings sit on their own rows above the lines they cover")


def two_tables_differ():
    rnd = random.Random(42)
    rows = [["Table A. Enterprises by country"], ["country", "enterprises", "employees"]]
    spots = []
    for c in EU[:6]:
        e, p = rnd.randint(1000, 90000), rnd.randint(10000, 900000)
        rows.append([c, e, p])
        if c == "Denmark":
            spots.append({"country": c, "enterprises": e, "employees": p})
    rows += [[], ["Table B. Enterprises by size class"], ["size_class", "share_pct"]]
    for sc in ["10-49", "50-249", "250+"]:
        rows.append([sc, round(rnd.uniform(1, 80), 1)])
    write_text("two_different_tables.csv", csv_text(rows))
    expect("two_different_tables.csv", [["country", "category"], ["enterprises", "number"], ["employees", "number"]], 6, spots,
           "Two unrelated tables one under the other. The first is read and the second is offered as an alternative")


def names_and_units():
    rnd = random.Random(43)
    rows = [["Region", "GDP", "Population", "R&D spending"], ["", "EUR bn", "thousand", "% of GDP"]]
    spots = []
    for r in ["Brussels", "Flanders", "Wallonia"]:
        vals = [round(rnd.uniform(50, 300), 1), rnd.randint(1000, 7000), round(rnd.uniform(1, 4), 2)]
        rows.append([r] + vals)
        if r == "Flanders":
            spots.append(dict(zip(["Region", "GDP EUR bn", "Population thousand", "R&D spending % of GDP"], [r] + vals)))
    write_text("names_with_unit_row.csv", csv_text(rows))
    expect("names_with_unit_row.csv", [["Region", "category"], ["GDP EUR bn", "number"], ["Population thousand", "number"],
                                       ["R&D spending % of GDP", "number"]], 3, spots, "Column names on one row and their units on the next")


def main():
    for f in HERE.iterdir():
        if f.name.startswith(("build_", "__")) or f.name == "README.md":
            continue
        f.unlink()
    for fn in [stress, eurostat_single, eurostat_noflags, eurostat_tsv, eurostat_tsv_size, eurostat_sdmx_csv, oecd_stat, oecd_explorer,
               oecd_wide, wdi, wdi_xlsx, ilostat, statbel, genesis, french_export, german_export, excel_report, excel_serials,
               crosstab_csv, crosstab_xlsx, clean_csv, clean_xlsx, json_records, json_stat, bom_csv, latin1_csv, crlf_csv, wide60,
               long50k, repeated_header, stacked_csv, text_only, quarters_wide, lightcast, nace_sum, pandas_index, merged_rows,
               pipe_csv, footnote_header_csv, side_by_side, no_header, group_rows, two_tables_differ, names_and_units]:
        fn()


if __name__ == "__main__":
    main()
